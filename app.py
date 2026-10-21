import io
import streamlit as st
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

# Configuration de la page Streamlit
st.set_page_config(
    page_title="Générateur de support de présentation d'un scénario Pédagogique EPS & numérique",
    page_icon="⚡",
    layout="wide",
)

# Style CSS personnalisé pour coller à l'ambiance sombre / jaune fluo
st.markdown(
    """
    <style>
    .main {
        background-color: #09090b;
        color: #f4f4f5;
    }
    h1, h2, h3 {
        color: #ccff00 !important;
        font-family: 'Calibri', sans-serif;
    }
    .stButton>button {
        background-color: #ccff00;
        color: #09090b;
        font-weight: bold;
        border: none;
        padding: 0.5rem 1rem;
        border-radius: 4px;
    }
    .stButton>button:hover {
        background-color: #b3ff00;
        color: #000000;
    }
    </style>
""",
    unsafe_allow_html=True,
)

st.title("⚡ Générateur de support de présentation d'un scénario Pédagogique EPS & numérique")
st.markdown(
    "Complétez les rubriques ci-dessous afin de générer votre présentation en conformité avec la charte graphique de référence. Vous pourrez ensuite télécharger un Powerpoint que vous pourrez modifier afin d'y ajouter ou soustraire des éléments."
)

with st.form("scenario_form"):
    # --- PARTIE 01 : CADRAGE ---
    st.header("01 / CADRAGE : Identité & Contexte")
    col1, col2 = st.columns(2)
    with col1:
        titre_apsa = st.text_input(
            "Titre du projet & APSA",
            placeholder="ex: Gymnastique - La roulade avant",
        )
        niveau_ca = st.text_input(
            "Niveau & Champ d'Apprentissage (CA)", placeholder="ex: Seconde - CA1"
        )
    with col2:
        auteurs = st.text_input(
            "Auteurs & Établissement",
            placeholder="ex: Nom des collègues - Lycée X",
        )
        constat_initial = st.text_area(
            "Constat initial (Obstacle d'apprentissage)",
            placeholder=(
                "Quel obstacle d'apprentissage a motivé l'introduction du"
                " numérique ?"
            ),
        )

    # --- PARTIE 02 : CIBLES ---
    st.header("02 / CIBLES : Transformations Visées chez l'Élève")
    focus_eleve = st.text_input(
        "Focus principal",
        placeholder=(
            "L'élève doit passer d'une perception subjective à une analyse"
            " objective..."
        ),
    )
    col3, col4, col5 = st.columns(3)
    with col3:
        dim_motrice = st.text_area(
            "Dimension Motrice",
            placeholder="Gains de coordination, réajustement des postures...",
        )
    with col4:
        methodes_analyse = st.text_area(
            "Méthodes & Analyse",
            placeholder="Capacité à décoder son propre mouvement...",
        )
    with col5:
        roles_sociaux = st.text_area(
            "Rôles Sociaux",
            placeholder="Engagement actif dans le rôle d'observateur, de coach...",
        )

    # --- PARTIE 03 : DISPOSITIF ---
    st.header("03 / DISPOSITIF : Le Choix Numérique")
    col6, col7 = st.columns(2)
    with col6:
        outil_appli = st.text_input(
            "Outil & Application",
            placeholder="ex: BaM Video Delay, Genially, tablettes...",
        )
        plus_value = st.text_area(
            "Plus-value Pédagogique",
            placeholder="En quoi l'outil fait-il ce que le papier ne permet pas ?",
        )
    with col7:
        temps_pratique = st.text_area(
            "Temps de Pratique Moteur",
            placeholder=(
                "Organisation spatiale pensée pour maintenir un taux d'effort"
                " élevé."
            ),
        )

    # --- PARTIE 04 : SÉQUENCE ---
    st.header("04 / SÉQUENCE : Mise en Œuvre en Leçon")

    st.subheader("Phase 01 (5-10 min) : Lancement")
    phase1_desc = st.text_input(
        "Description Lancement",
        value=(
            "Mise en train, installation sécurisée des ateliers et rappel des"
            " consignes numériques."
        ),
    )

    st.subheader("Phase 02 (30-40 min) : Situation d'Apprentissage Principale")
    col8, col9 = st.columns(2)
    with col8:
        p2_dispositif = st.text_area(
            "Dispositif", placeholder="Organisation spatiale & ateliers."
        )
        p2_realisation = st.text_area(
            "Critères de Réalisation", placeholder="Consignes d'exécution du geste."
        )
        p2_reussite = st.text_area(
            "Critères de Réussite", placeholder="Indicateurs pour valider la tâche."
        )
    with col9:
        p2_regulations = st.text_area(
            "Régulations Enseignant", placeholder="Conseils et ajustements ciblés."
        )
        p2_roles_moteur = st.text_area(
            "Rôles Élèves (Activité motrice)",
            placeholder="Réalisation et répétitions.",
        )
        p2_roles_num = st.text_area(
            "Rôles Élèves (Outil numérique)",
            placeholder="Saisie, observation vidéo, co-évaluation.",
        )

    st.subheader("Phase 03 (5-10 min) : Bilan & Trace")
    phase3_desc = st.text_input(
        "Description Bilan",
        value=(
            "Verbalisation des progrès, export des données collectives ou"
            " sauvegarde des acquis."
        ),
    )

    # --- PARTIE 05 : RETOUR D'EXPÉRIENCE ---
    st.header("05 / RETOUR D'EXPÉRIENCE : Vigilance & Conseils")
    col10, col11, col12 = st.columns(3)
    with col10:
        obs_tech = st.text_area(
            "Obstacles Techniques",
            placeholder=(
                "Batteries, luminosité, réseau absent : prévoir solution secours."
            ),
        )
    with col11:
        derive_chrono = st.text_area(
            "Dérive Chronophage",
            placeholder=(
                "Temps excessif devant les écrans. Fixer des règles (max 30s)."
            ),
        )
    with col12:
        cadre_ethique = st.text_area(
            "Cadre Éthique & RGPD",
            placeholder=(
                "Stockage local uniquement, suppression des vidéos, droit à"
                " l'image."
            ),
        )

    submitted = st.form_submit_button(
        "Générer le PowerPoint de référence (.pptx)"
    )


# --- LOGIQUE DE GÉNÉRATION DU PPTX EN 5 DIAPOS AVEC CARTES ---
def create_presentation(data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    BG_COLOR = RGBColor(9, 9, 11)        # #09090b (Noir minéral)
    CARD_BG = RGBColor(24, 24, 27)       # #18181b (Gris très foncé pour les cartes)
    ACCENT_COLOR = RGBColor(204, 255, 0) # #ccff00 (Jaune fluo)
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(161, 161, 170)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category_text="SCÉNARIO PÉDAGOGIQUE EPS"):
        set_slide_background(slide)
        
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_COLOR
        p_cat.font.name = "Calibri"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Calibri"

    def add_card(slide, left, top, width, height, title, content_list):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = ACCENT_COLOR
        shape.line.width = Pt(1)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_COLOR
        p.font.name = "Calibri"
        p.space_after = Pt(10)

        for item in content_list:
            p_item = tf.add_paragraph()
            # Si l'élément a déjà un libellé formaté, on l'affiche proprement
            p_item.text = f"• {item}" if not item.startswith("•") else item
            p_item.font.size = Pt(12)
            p_item.font.color.rgb = TEXT_WHITE
            p_item.font.name = "Calibri"
            p_item.space_after = Pt(6)

    # ==========================================
    # DIAPOSITIVE 1 : PAGE DE GARDE (ACCUEIL)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    title_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.8), Inches(10.33), Inches(4.0))
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = CARD_BG
    title_shape.line.color.rgb = ACCENT_COLOR
    title_shape.line.width = Pt(1.5)

    tf_home = title_shape.text_frame
    tf_home.word_wrap = True
    tf_home.margin_left = Inches(0.5)
    tf_home.margin_right = Inches(0.5)
    tf_home.margin_top = Inches(0.8)

    p_badge = tf_home.paragraphs[0]
    p_badge.text = "GROUPE ACADÉMIQUE EPS"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_COLOR
    p_badge.font.name = "Calibri"
    p_badge.space_after = Pt(15)

    p_main = tf_home.add_paragraph()
    p_main.text = f"Scénario Pédagogique\n{data.get('titre_apsa', '')}"
    p_main.font.size = Pt(32)
    p_main.font.bold = True
    p_main.font.color.rgb = TEXT_WHITE
    p_main.font.name = "Calibri"
    p_main.space_after = Pt(15)

    p_sub = tf_home.add_paragraph()
    p_sub.text = f"Niveau / CA : {data.get('niveau_ca', '')}  |  Auteurs : {data.get('auteurs', '')}"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.font.name = "Calibri"

    # ==========================================
    # DIAPOSITIVE 2 : CIBLES
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Cibles - Transformations Visées", "02. Cibles")

    add_card(slide2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "🎯 Focus & Constat Initial", 
             [f"Focus principal : {data.get('focus_eleve', '')}",
              f"Constat initial : {data.get('constat_initial', '')}"])

    add_card(slide2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), 
             "⚙️ Dimensions & Rôles", 
             [f"Dimension motrice : {data.get('dim_motrice', '')}",
              f"Méthodes & Analyse : {data.get('methodes_analyse', '')}",
              f"Rôles sociaux : {data.get('roles_sociaux', '')}"])

    # ==========================================
    # DIAPOSITIVE 3 : DISPOSITIF
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Dispositif - Le Choix Numérique", "03. Dispositif")

    add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "💻 Outil & Plus-Value", 
             [f"Outil & Application : {data.get('outil_appli', '')}",
              f"Plus-value pédagogique : {data.get('plus_value', '')}"])

    add_card(slide3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), 
             "⏱️ Pratique & Règle d'or", 
             [f"Temps de pratique : {data.get('temps_pratique', '')}",
              "⚠️ Règle d'or : L'écran ne doit jamais remplacer la pratique physique."])

    # ==========================================
    # DIAPOSITIVE 4 : SÉQUENCE (3 PHASES)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Séquence - Mise en Œuvre en Leçon", "04. Séquence")

    add_card(slide4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "Phase 1 (5-10 min) : Lancement", 
             [data.get('phase1_desc', '')])

    add_card(slide4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), 
             "Phase 2 (30-40 min) : Apprentissage", 
             [f"Dispositif : {data.get('p2_dispositif', '')}",
              f"Réalisation : {data.get('p2_realisation', '')}",
              f"Réussite : {data.get('p2_reussite', '')}",
              f"Régulations : {data.get('p2_regulations', '')}",
              f"Rôles moteur : {data.get('p2_roles_moteur', '')}",
              f"Rôles num. : {data.get('p2_roles_num', '')}"])

    add_card(slide4, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0), 
             "Phase 3 (5-10 min) : Bilan & Trace", 
             [data.get('phase3_desc', '')])

    # ==========================================
    # DIAPOSITIVE 5 : RETOUR D'EXPÉRIENCE
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Retour d'Expérience - Vigilance & Conseils", "05. Retour d'Expérience")

    add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), 
             "⚠️ Obstacles & Dérives", 
             [f"Obstacles techniques : {data.get('obs_tech', '')}",
              f"Dérive chronophage : {data.get('derive_chrono', '')}"])

    add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), 
             "🛡️ Cadre Éthique & RGPD", 
             [f"Cadre éthique : {data.get('cadre_ethique', '')}",
              "💡 Une fiche réussie permet à un collègue de reproduire le scénario."])

    # Sauvegarde en mémoire binaire
    ppt_io = io.BytesIO()
    prs.save(ppt_io)
    ppt_io.seek(0)
    return ppt_io


# Exécution du bouton de téléchargement
if submitted:
    data_dict = {
        "titre_apsa": titre_apsa,
        "niveau_ca": niveau_ca,
        "auteurs": auteurs,
        "constat_initial": constat_initial,
        "focus_eleve": focus_eleve,
        "dim_motrice": dim_motrice,
        "methodes_analyse": methodes_analyse,
        "roles_sociaux": roles_sociaux,
        "outil_appli": outil_appli,
        "plus_value": plus_value,
        "temps_pratique": temps_pratique,
        "phase1_desc": phase1_desc,
        "p2_dispositif": p2_dispositif,
        "p2_realisation": p2_realisation,
        "p2_reussite": p2_reussite,
        "p2_regulations": p2_regulations,
        "p2_roles_moteur": p2_roles_moteur,
        "p2_roles_num": p2_roles_num,
        "phase3_desc": phase3_desc,
        "obs_tech": obs_tech,
        "derive_chrono": derive_chrono,
        "cadre_ethique": cadre_ethique,
    }

    ppt_file = create_presentation(data_dict)

    st.success(
        "🎉 Votre scénario pédagogique a été généré avec succès selon la charte"
        " graphique de référence !"
    )
    st.download_button(
        label="📥 Télécharger le fichier PowerPoint (.pptx)",
        data=ppt_file,
        file_name="Scenario_Pedagogique_EPS.pptx",
        mime=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
    )